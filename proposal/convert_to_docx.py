from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
import re

md_path = "proposal/Proton_Integration_Playwright_Proposal.md"
docx_path = "proposal/Proton_Integration_Playwright_Proposal.docx"

doc = Document()

# Page margins
for section in doc.sections:
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1.2)
    section.right_margin = Inches(1.2)

def set_heading(para, level):
    sizes = {1: 22, 2: 16, 3: 13}
    colors = {1: RGBColor(0x1F, 0x49, 0x7D), 2: RGBColor(0x2E, 0x74, 0xB5), 3: RGBColor(0x1F, 0x49, 0x7D)}
    for run in para.runs:
        run.font.size = Pt(sizes.get(level, 12))
        run.font.bold = True
        run.font.color.rgb = colors.get(level, RGBColor(0, 0, 0))

with open(md_path, "r") as f:
    lines = f.readlines()

i = 0
while i < len(lines):
    line = lines[i].rstrip()

    # Skip horizontal rules
    if re.match(r'^---+$', line):
        i += 1
        continue

    # Headings
    m = re.match(r'^(#{1,3})\s+(.*)', line)
    if m:
        level = len(m.group(1))
        text = re.sub(r'\*\*(.*?)\*\*', r'\1', m.group(2))
        text = re.sub(r'`(.*?)`', r'\1', text)
        para = doc.add_heading(text, level=level)
        set_heading(para, level)
        i += 1
        continue

    # Blockquote
    if line.startswith('>'):
        text = line.lstrip('> ').strip()
        text = re.sub(r'\*\*(.*?)\*\*', r'\1', text)
        para = doc.add_paragraph(text)
        para.paragraph_format.left_indent = Inches(0.4)
        for run in para.runs:
            run.font.italic = True
            run.font.color.rgb = RGBColor(0x60, 0x60, 0x60)
        i += 1
        continue

    # Table
    if line.startswith('|'):
        table_lines = []
        while i < len(lines) and lines[i].strip().startswith('|'):
            table_lines.append(lines[i].strip())
            i += 1
        # Filter separator rows
        rows = [r for r in table_lines if not re.match(r'^\|[\s\-\|:]+\|$', r)]
        if not rows:
            continue
        cells_per_row = [re.split(r'\s*\|\s*', r.strip('|').strip()) for r in rows]
        col_count = max(len(c) for c in cells_per_row)
        table = doc.add_table(rows=len(cells_per_row), cols=col_count)
        table.style = 'Table Grid'
        for ri, row_cells in enumerate(cells_per_row):
            for ci, cell_text in enumerate(row_cells):
                if ci >= col_count:
                    break
                cell_text = re.sub(r'\*\*(.*?)\*\*', r'\1', cell_text)
                cell_text = re.sub(r'`(.*?)`', r'\1', cell_text)
                cell = table.cell(ri, ci)
                cell.text = cell_text.strip()
                if ri == 0:
                    for run in cell.paragraphs[0].runs:
                        run.font.bold = True
                        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                    cell._tc.get_or_add_tcPr()
                    from docx.oxml.ns import qn
                    from docx.oxml import OxmlElement
                    shd = OxmlElement('w:shd')
                    shd.set(qn('w:fill'), '2E74B5')
                    shd.set(qn('w:color'), 'auto')
                    shd.set(qn('w:val'), 'clear')
                    cell._tc.tcPr.append(shd)
        doc.add_paragraph()
        continue

    # Code block
    if line.startswith('```'):
        code_lines = []
        i += 1
        while i < len(lines) and not lines[i].strip().startswith('```'):
            code_lines.append(lines[i].rstrip())
            i += 1
        i += 1  # skip closing ```
        para = doc.add_paragraph('\n'.join(code_lines))
        para.paragraph_format.left_indent = Inches(0.4)
        for run in para.runs:
            run.font.name = 'Courier New'
            run.font.size = Pt(8)
            run.font.color.rgb = RGBColor(0x20, 0x20, 0x20)
        continue

    # Bullet list
    if re.match(r'^[-*]\s+', line):
        text = re.sub(r'^[-*]\s+', '', line)
        text = re.sub(r'\*\*(.*?)\*\*', r'\1', text)
        text = re.sub(r'`(.*?)`', r'\1', text)
        doc.add_paragraph(text, style='List Bullet')
        i += 1
        continue

    # Empty line
    if not line.strip():
        i += 1
        continue

    # Normal paragraph
    text = re.sub(r'\*\*(.*?)\*\*', r'\1', line)
    text = re.sub(r'`(.*?)`', r'\1', text)
    doc.add_paragraph(text)
    i += 1

doc.save(docx_path)
print(f"Saved: {docx_path}")
